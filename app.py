"""
Equation-Sie — Outil de prospection géolocalisée
Croise la base Pipedrive (organisations + contacts) avec une adresse pivot
et un rayon, pour sortir les sociétés et contacts dans la zone.
"""

import os
import json
import time
import io
from datetime import datetime, timezone
from math import radians, cos, sin, asin, sqrt

import streamlit as st
import pandas as pd
import requests

# ─────────────────────────────────────────────────────────────────────────────
# Configuration
# ─────────────────────────────────────────────────────────────────────────────
DATA_DIR = "/var/data"
ORGS_FILE = os.path.join(DATA_DIR, "organisations.xlsx")
CONTACTS_FILE = os.path.join(DATA_DIR, "contacts.xlsx")
GEOCODED_FILE = os.path.join(DATA_DIR, "organisations_geocodees.xlsx")
META_FILE = os.path.join(DATA_DIR, "metadata.json")
HISTORY_FILE = os.path.join(DATA_DIR, "history.json")
HISTORY_MAX = 20  # Nombre max d'entrées conservées

BAN_SEARCH = "https://api-adresse.data.gouv.fr/search/"
BAN_BATCH = "https://api-adresse.data.gouv.fr/search/csv/"

# Mot de passe admin (à changer dans les Secrets HF : ADMIN_PASSWORD)
# .strip() pour nettoyer les sauts de ligne / espaces parasites du copier-coller
ADMIN_PASSWORD = os.environ.get("ADMIN_PASSWORD", "").strip()

# Mot de passe d'accès à l'application (commerciaux)
# Si non configuré, l'app est ouverte (à éviter en production)
APP_PASSWORD = os.environ.get("APP_PASSWORD", "").strip()

st.set_page_config(
    page_title="Equation-Sie — Prospection géolocalisée",
    page_icon="📍",
    layout="wide",
)

# ─────────────────────────────────────────────────────────────────────────────
# Métadonnées (dates d'import)
# ─────────────────────────────────────────────────────────────────────────────
def load_metadata():
    if os.path.exists(META_FILE):
        with open(META_FILE, "r") as f:
            return json.load(f)
    return {
        "orgs_uploaded_at": None,
        "orgs_rows": 0,
        "contacts_uploaded_at": None,
        "contacts_rows": 0,
        "geocoded_at": None,
        "geocoded_rows": 0,
        "geocoded_success": 0,
    }


def save_metadata(meta):
    os.makedirs(DATA_DIR, exist_ok=True)
    with open(META_FILE, "w") as f:
        json.dump(meta, f, indent=2)


def fmt_date(iso_str):
    if not iso_str:
        return "—"
    dt = datetime.fromisoformat(iso_str)
    return dt.strftime("%d/%m/%Y %H:%M")


def freshness_status(iso_str):
    """🟢 < 30j, 🟠 30-60j, 🔴 > 60j"""
    if not iso_str:
        return "⚪ Aucun fichier"
    dt = datetime.fromisoformat(iso_str)
    age_days = (datetime.now(timezone.utc) - dt).days
    if age_days < 30:
        return f"🟢 À jour ({age_days} j)"
    elif age_days < 60:
        return f"🟠 À rafraîchir ({age_days} j)"
    return f"🔴 Périmé ({age_days} j)"


# ─────────────────────────────────────────────────────────────────────────────
# Historique des recherches
# ─────────────────────────────────────────────────────────────────────────────
def load_history():
    if os.path.exists(HISTORY_FILE):
        try:
            with open(HISTORY_FILE, "r") as f:
                return json.load(f)
        except Exception:
            return []
    return []


def save_history(history):
    os.makedirs(DATA_DIR, exist_ok=True)
    with open(HISTORY_FILE, "w") as f:
        json.dump(history, f, indent=2, ensure_ascii=False)


def add_to_history(pivot_address, radius_m, label, nb_sociétés, nb_avec_contact):
    """Ajoute une entrée à l'historique. Dédoublonne sur (adresse + rayon)."""
    history = load_history()
    new_entry = {
        "timestamp": datetime.now(timezone.utc).isoformat(),
        "pivot_address": pivot_address,
        "pivot_label": label,
        "radius_m": radius_m,
        "nb_societes": nb_sociétés,
        "nb_avec_contact": nb_avec_contact,
    }
    # Supprimer les doublons exacts (même adresse + même rayon)
    history = [
        h for h in history
        if not (h.get("pivot_address") == pivot_address and h.get("radius_m") == radius_m)
    ]
    # Insérer en tête
    history.insert(0, new_entry)
    # Tronquer
    history = history[:HISTORY_MAX]
    save_history(history)


def clear_history():
    if os.path.exists(HISTORY_FILE):
        os.remove(HISTORY_FILE)


def fmt_relative(iso_str):
    """Formate un timestamp en 'il y a X' relatif."""
    dt = datetime.fromisoformat(iso_str)
    delta = datetime.now(timezone.utc) - dt
    seconds = delta.total_seconds()
    if seconds < 60:
        return "à l'instant"
    if seconds < 3600:
        return f"il y a {int(seconds//60)} min"
    if seconds < 86400:
        return f"il y a {int(seconds//3600)} h"
    if seconds < 86400 * 7:
        return f"il y a {int(seconds//86400)} j"
    return dt.strftime("%d/%m/%Y")



# ─────────────────────────────────────────────────────────────────────────────
# Géocodage BAN
# ─────────────────────────────────────────────────────────────────────────────
def geocode_address(address: str):
    """Géocode une adresse unique via l'API BAN. Retourne (lat, lon, label, score)."""
    for attempt in range(4):
        try:
            time.sleep(2 * attempt)
            r = requests.get(BAN_SEARCH, params={"q": address, "limit": 1}, timeout=15)
            if r.status_code == 200:
                data = r.json()
                if data.get("features"):
                    feat = data["features"][0]
                    lon, lat = feat["geometry"]["coordinates"]
                    props = feat["properties"]
                    return lat, lon, props["label"], props["score"]
        except Exception:
            pass
    return None, None, None, None


def geocode_batch(df_orgs, progress_bar=None):
    """Géocode en batch via l'API BAN. df_orgs doit contenir id + adresse_clean."""
    csv_buffer = io.StringIO()
    df_orgs[["id", "adresse_clean"]].to_csv(csv_buffer, index=False)
    csv_bytes = csv_buffer.getvalue().encode("utf-8")

    for attempt in range(4):
        try:
            if progress_bar:
                progress_bar.progress(0.1, text=f"Tentative {attempt+1}/4 — Envoi à la BAN…")
            time.sleep(5 * attempt)
            r = requests.post(
                BAN_BATCH,
                files={"data": ("orgs.csv", csv_bytes, "text/csv")},
                data={"columns": "adresse_clean"},
                timeout=1800,
            )
            if r.status_code == 200:
                if progress_bar:
                    progress_bar.progress(0.9, text="Réponse reçue, traitement…")
                return pd.read_csv(io.BytesIO(r.content), dtype={"id": str})
        except Exception as e:
            if progress_bar:
                progress_bar.progress(0.1, text=f"Tentative {attempt+1} échec : {e}")
    raise RuntimeError("Échec du géocodage batch après 4 tentatives")


def haversine(lat1, lon1, lat2, lon2):
    R = 6371000
    lat1, lon1, lat2, lon2 = map(radians, [lat1, lon1, lat2, lon2])
    dlat = lat2 - lat1
    dlon = lon2 - lon1
    a = sin(dlat / 2) ** 2 + cos(lat1) * cos(lat2) * sin(dlon / 2) ** 2
    return 2 * R * asin(sqrt(a))


# ─────────────────────────────────────────────────────────────────────────────
# Pipeline complet d'enrichissement
# ─────────────────────────────────────────────────────────────────────────────
def process_orgs_file(file_bytes, progress_bar=None):
    """Lit, géocode et sauvegarde le fichier organisations."""
    df = pd.read_excel(io.BytesIO(file_bytes))
    df["Organisation - ID"] = df["Organisation - ID"].astype(str)

    if progress_bar:
        progress_bar.progress(0.05, text=f"{len(df)} lignes chargées — Préparation…")

    # Préparer pour la BAN
    df_to_geo = df[df["Organisation - Adresse"].notna()].copy()
    df_to_geo["adresse_clean"] = (
        df_to_geo["Organisation - Adresse"]
        .str.replace(", France", "", regex=False)
        .str.replace(",France", "", regex=False)
        .str.strip()
    )
    geo_input = df_to_geo[["Organisation - ID", "adresse_clean"]].copy()
    geo_input.columns = ["id", "adresse_clean"]

    # Lancer le batch
    geo_result = geocode_batch(geo_input, progress_bar=progress_bar)

    # Fusionner
    geo_keep = geo_result[
        ["id", "latitude", "longitude", "result_score", "result_label",
         "result_postcode", "result_city", "result_type"]
    ].copy()
    geo_keep.columns = [
        "Organisation - ID", "Latitude", "Longitude", "Score_geocodage",
        "Adresse_normalisee", "Code_postal", "Ville", "Type_matching",
    ]
    merged = df.merge(geo_keep, on="Organisation - ID", how="left")

    if progress_bar:
        progress_bar.progress(0.95, text="Sauvegarde…")

    os.makedirs(DATA_DIR, exist_ok=True)
    merged.to_excel(GEOCODED_FILE, index=False)
    # On garde aussi le fichier brut pour traçabilité
    with open(ORGS_FILE, "wb") as f:
        f.write(file_bytes)

    return len(merged), int(merged["Latitude"].notna().sum())


def save_contacts_file(file_bytes):
    df = pd.read_excel(io.BytesIO(file_bytes))
    os.makedirs(DATA_DIR, exist_ok=True)
    with open(CONTACTS_FILE, "wb") as f:
        f.write(file_bytes)
    return len(df)


# ─────────────────────────────────────────────────────────────────────────────
# Requête de proximité
# ─────────────────────────────────────────────────────────────────────────────
@st.cache_data(show_spinner=False)
def load_geocoded():
    if not os.path.exists(GEOCODED_FILE):
        return None
    return pd.read_excel(GEOCODED_FILE)


@st.cache_data(show_spinner=False)
def load_contacts():
    if not os.path.exists(CONTACTS_FILE):
        return None
    return pd.read_excel(CONTACTS_FILE)


def run_proximity_search(pivot_address: str, radius_m: int):
    lat_p, lon_p, label, score = geocode_address(pivot_address)
    if lat_p is None:
        return None, "Impossible de géocoder l'adresse pivot."

    orgs = load_geocoded()
    contacts = load_contacts()
    if orgs is None:
        return None, "Aucun fichier organisations géocodé. Va dans l'onglet Admin pour en charger un."
    if contacts is None:
        return None, "Aucun fichier contacts. Va dans l'onglet Admin pour en charger un."

    orgs = orgs.dropna(subset=["Latitude", "Longitude"]).copy()
    orgs["Organisation - ID"] = orgs["Organisation - ID"].astype(str)
    orgs["Distance_m"] = orgs.apply(
        lambda r: haversine(lat_p, lon_p, r["Latitude"], r["Longitude"]), axis=1
    )
    zone = orgs[orgs["Distance_m"] <= radius_m].copy()
    zone["Distance_m"] = zone["Distance_m"].round(0).astype(int)
    zone_ids = set(zone["Organisation - ID"])

    contacts = contacts.dropna(subset=["Organisation - ID"]).copy()
    contacts["Organisation - ID"] = contacts["Organisation - ID"].astype(int).astype(str)
    contacts_zone = contacts[contacts["Organisation - ID"].isin(zone_ids)].copy()

    phone_cols = [c for c in [
        "Personne - PORTABLE", "Personne - Téléphone - Travail",
        "Personne - Téléphone - Domicile", "Personne - Téléphone - Mobile",
        "Personne - Téléphone - Autre"
    ] if c in contacts_zone.columns]
    if phone_cols:
        contacts_zone["has_phone"] = contacts_zone[phone_cols].notna().any(axis=1)
    else:
        contacts_zone["has_phone"] = False

    orgs_avec = set(contacts_zone["Organisation - ID"])
    orgs_sans = zone_ids - orgs_avec

    return {
        "pivot_label": label,
        "pivot_score": score,
        "pivot_lat": lat_p,
        "pivot_lon": lon_p,
        "radius_m": radius_m,
        "zone": zone,
        "contacts_zone": contacts_zone,
        "orgs_avec_ids": orgs_avec,
        "orgs_sans_ids": orgs_sans,
    }, None


def build_export_xlsx(result):
    """Construit le fichier Excel multi-onglets en mémoire."""
    zone = result["zone"]
    cz = result["contacts_zone"]

    # ─── Préparation : lookup enrichi avec TOUTES les colonnes métier orgs ──
    # Colonnes métier à ramener depuis la base organisations
    orgs_meta_cols = [
        "Organisation - ID", "Organisation - Nom", "Distance_m",
        "Organisation - Adresse", "Organisation - ADRESSE VERIFIEE",
        "Organisation - NBR DE POSTES", "Organisation - SURFACE OCCUPEE",
        "Organisation - TYPE DE CONTRAT", "Organisation - OPERATEUR FLEX",
        "Organisation - ANNEE DEBUT DE BAIL",
        "Organisation - PROCHAINE ECHEANCE DE CONTRAT",
        "Organisation - Date de fin de PROCHAINE ECHEANCE DE CONTRAT",
        "Organisation - Profil LinkedIn",
    ]
    orgs_meta_cols = [c for c in orgs_meta_cols if c in zone.columns]
    zone_lookup = zone[orgs_meta_cols].copy()

    # Renommage des colonnes métier pour usage commun (onglets 1 et 2)
    meta_rename = {
        "Organisation - Nom": "Société",
        "Distance_m": "Distance (m)",
        "Organisation - Adresse": "Adresse société",
        "Organisation - ADRESSE VERIFIEE": "Adresse vérifiée",
        "Organisation - NBR DE POSTES": "Nb postes",
        "Organisation - SURFACE OCCUPEE": "Surface (m²)",
        "Organisation - TYPE DE CONTRAT": "Type contrat",
        "Organisation - OPERATEUR FLEX": "Opérateur flex",
        "Organisation - ANNEE DEBUT DE BAIL": "Année début bail",
        "Organisation - PROCHAINE ECHEANCE DE CONTRAT": "Prochaine échéance",
        "Organisation - Date de fin de PROCHAINE ECHEANCE DE CONTRAT": "Date fin échéance",
        "Organisation - Profil LinkedIn": "LinkedIn société",
    }
    zone_lookup = zone_lookup.rename(columns=meta_rename)

    # On enrichit cz avec les colonnes orgs (jointure sur Organisation - ID)
    cz_enriched = cz.merge(zone_lookup, on="Organisation - ID", how="left")

    # ─── Onglet 1 — Contacts (détail) avec infos société ──────────────────
    cols_contacts = [
        # Repères en tête : distance + identification société
        "Distance (m)", "Société",
        # 👤 BLOC CONTACT
        "Personne - Civilité",
        "Personne - Prénom", "Personne - Nom de famille", "Personne - FONCTION INTITULE",
        "Personne - E-mail - Travail", "Personne - E-mail - Domicile", "Personne - E-mail - Autre",
        "Personne - PORTABLE", "Personne - Téléphone - Travail",
        "Personne - Téléphone - Mobile", "Personne - Téléphone - Domicile", "Personne - Téléphone - Autre",
        "Personne - URL LINK", "Personne - STATUT LINKEDIN",
        # 🏢 BLOC SOCIÉTÉ
        "Adresse société", "Adresse vérifiée",
        "Nb postes", "Surface (m²)", "Type contrat", "Opérateur flex",
        "Année début bail", "Prochaine échéance", "Date fin échéance",
        "LinkedIn société",
        # IDs en fin
        "Personne - ID", "Organisation - ID",
    ]
    cols_contacts = [c for c in cols_contacts if c in cz_enriched.columns]
    contacts_export = cz_enriched[cols_contacts].copy()
    rename_map = {
        "Personne - Civilité": "Civilité",
        "Personne - Prénom": "Prénom",
        "Personne - Nom de famille": "Nom",
        "Personne - FONCTION INTITULE": "Fonction",
        "Personne - E-mail - Travail": "Email travail",
        "Personne - E-mail - Domicile": "Email domicile",
        "Personne - E-mail - Autre": "Email autre",
        "Personne - PORTABLE": "Portable",
        "Personne - Téléphone - Travail": "Tel travail",
        "Personne - Téléphone - Mobile": "Tel mobile",
        "Personne - Téléphone - Domicile": "Tel domicile",
        "Personne - Téléphone - Autre": "Tel autre",
        "Personne - URL LINK": "LinkedIn",
        "Personne - STATUT LINKEDIN": "Statut LinkedIn",
        "Personne - ID": "ID Personne",
        "Organisation - ID": "ID Org",
    }
    contacts_export = contacts_export.rename(columns=rename_map)
    sort_cols = [c for c in ["Distance (m)", "Société", "Nom"] if c in contacts_export.columns]
    contacts_export = contacts_export.sort_values(sort_cols)

    # ─── Onglet 2 — Récap sociétés AVEC contacts (+ infos métier) ─────────
    # On agrège les contacts par société, puis on jointe les infos métier
    recap_agg = cz.groupby(["Organisation - ID"]).agg(
        Nb_contacts=("Personne - ID", "count"),
        Nb_emails=("Personne - E-mail - Travail", lambda s: s.notna().sum()),
        Nb_telephones=("has_phone", "sum"),
        Nb_linkedin=("Personne - URL LINK", lambda s: s.notna().sum()),
    ).reset_index()

    # On jointe avec les infos métier
    recap = zone_lookup.merge(recap_agg, on="Organisation - ID", how="inner")

    # Renommer les colonnes agrégées
    recap = recap.rename(columns={
        "Organisation - ID": "ID Pipedrive",
        "Nb_contacts": "Nb contacts",
        "Nb_emails": "Nb emails",
        "Nb_telephones": "Nb tél.",
        "Nb_linkedin": "Nb LinkedIn",
    })

    # Ordre des colonnes : distance + société + stats contacts + métier
    recap_cols_order = [
        "Distance (m)", "Société", "Adresse société", "Adresse vérifiée",
        "Nb contacts", "Nb emails", "Nb tél.", "Nb LinkedIn",
        "Nb postes", "Surface (m²)", "Type contrat", "Opérateur flex",
        "Année début bail", "Prochaine échéance", "Date fin échéance",
        "LinkedIn société", "ID Pipedrive",
    ]
    recap_cols_order = [c for c in recap_cols_order if c in recap.columns]
    recap = recap[recap_cols_order].sort_values(["Distance (m)", "Société"])

    # ─── Onglet 3 — Sans contact (déjà enrichi) ───────────────────────────
    sans = zone[zone["Organisation - ID"].isin(result["orgs_sans_ids"])].copy()
    cols_sans = [
        "Distance_m", "Organisation - Nom",
        "Organisation - Adresse", "Organisation - ADRESSE VERIFIEE",
        "Organisation - NBR DE POSTES", "Organisation - SURFACE OCCUPEE",
        "Organisation - TYPE DE CONTRAT", "Organisation - OPERATEUR FLEX",
        "Organisation - ANNEE DEBUT DE BAIL",
        "Organisation - PROCHAINE ECHEANCE DE CONTRAT",
        "Organisation - Date de fin de PROCHAINE ECHEANCE DE CONTRAT",
        "Organisation - Profil LinkedIn", "Organisation - ID",
    ]
    cols_sans = [c for c in cols_sans if c in sans.columns]
    sans_export = sans[cols_sans].copy()
    sans_rename = {
        "Distance_m": "Distance (m)",
        "Organisation - Nom": "Société",
        "Organisation - Adresse": "Adresse",
        "Organisation - ADRESSE VERIFIEE": "Adresse vérifiée",
        "Organisation - NBR DE POSTES": "Nb postes",
        "Organisation - SURFACE OCCUPEE": "Surface (m²)",
        "Organisation - TYPE DE CONTRAT": "Type contrat",
        "Organisation - OPERATEUR FLEX": "Opérateur flex",
        "Organisation - ANNEE DEBUT DE BAIL": "Année début bail",
        "Organisation - PROCHAINE ECHEANCE DE CONTRAT": "Prochaine échéance",
        "Organisation - Date de fin de PROCHAINE ECHEANCE DE CONTRAT": "Date fin échéance",
        "Organisation - Profil LinkedIn": "LinkedIn société",
        "Organisation - ID": "ID Pipedrive",
    }
    sans_export = sans_export.rename(columns=sans_rename).sort_values("Distance (m)")

    # ─── Onglet 0 — Synthèse (page de garde) ──────────────────────────────
    pivot_label = result.get("pivot_label", "—")
    radius_m = result.get("radius_m", "—")
    nb_zone = len(zone)
    nb_avec = len(result["orgs_avec_ids"])
    nb_sans = len(result["orgs_sans_ids"])
    nb_contacts_total = len(cz)
    date_recherche = datetime.now().strftime("%d/%m/%Y %H:%M")

    synthese = pd.DataFrame([
        ["📅 Date de la recherche", date_recherche],
        ["📍 Adresse pivot", pivot_label],
        ["📏 Rayon", f"{radius_m} m"],
        ["", ""],
        ["🏢 Sociétés dans la zone", nb_zone],
        ["✅ Avec contact", nb_avec],
        ["❌ Sans contact", nb_sans],
        ["👥 Contacts détaillés", nb_contacts_total],
    ], columns=["Indicateur", "Valeur"])

    # Écriture
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        synthese.to_excel(writer, sheet_name="Synthèse", index=False)
        contacts_export.to_excel(writer, sheet_name="Contacts (détail)", index=False)
        recap.to_excel(writer, sheet_name="Récap sociétés AVEC contacts", index=False)
        sans_export.to_excel(writer, sheet_name="Sociétés SANS contact", index=False)

        # Mise en forme de l'onglet Synthèse
        ws = writer.sheets["Synthèse"]
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        # Largeurs colonnes
        ws.column_dimensions["A"].width = 35
        ws.column_dimensions["B"].width = 45

        # Titre en haut (insertion ligne au-dessus)
        ws.insert_rows(1)
        ws["A1"] = "📍 Résultats de prospection géolocalisée"
        ws["A1"].font = Font(name="Calibri", size=16, bold=True, color="4A3B6B")
        ws.merge_cells("A1:B1")
        ws.row_dimensions[1].height = 30

        # Style en-têtes (ligne 2)
        header_fill = PatternFill(start_color="4A3B6B", end_color="4A3B6B", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        for cell in ws[2]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[2].height = 22

        # Style lignes de données
        thin_border = Border(
            left=Side(style="thin", color="DDDDDD"),
            right=Side(style="thin", color="DDDDDD"),
            top=Side(style="thin", color="DDDDDD"),
            bottom=Side(style="thin", color="DDDDDD"),
        )
        for row in ws.iter_rows(min_row=3, max_row=ws.max_row, min_col=1, max_col=2):
            for cell in row:
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
                cell.font = Font(name="Calibri", size=11)
            ws.row_dimensions[row[0].row].height = 22

        # Mise en gras de la colonne Indicateur
        for row in ws.iter_rows(min_row=3, max_row=ws.max_row, min_col=1, max_col=1):
            for cell in row:
                cell.font = Font(name="Calibri", size=11, bold=True)

    return output.getvalue()


# ─────────────────────────────────────────────────────────────────────────────
# UI
# ─────────────────────────────────────────────────────────────────────────────
st.title("📍 Equation-Sie — Prospection géolocalisée")
st.caption("Croise la base Pipedrive avec une adresse pivot pour sortir les sociétés et contacts dans la zone.")


# ─── Authentification globale (mot de passe d'équipe) ───────────────────────
def check_app_password():
    """
    Affiche un écran de login tant que l'utilisateur n'a pas saisi le bon mot de passe.
    Si APP_PASSWORD n'est pas configuré, l'app est ouverte (utile en dev).
    """
    if not APP_PASSWORD:
        return True  # Pas de mot de passe configuré, accès libre

    if st.session_state.get("authenticated"):
        return True

    # Écran de login centré
    _, col_login, _ = st.columns([1, 2, 1])
    with col_login:
        with st.container(border=True):
            st.markdown("### 🔐 Accès sécurisé")
            st.caption("Saisis le mot de passe d'équipe pour accéder à l'outil.")

            with st.form("login_form", clear_on_submit=False):
                pwd_input = st.text_input("Mot de passe", type="password", key="login_pwd")
                submit = st.form_submit_button("🔓 Se connecter", type="primary", use_container_width=True)

            if submit:
                if pwd_input == APP_PASSWORD:
                    st.session_state["authenticated"] = True
                    st.rerun()
                else:
                    st.error("❌ Mot de passe incorrect.")
    return False


if not check_app_password():
    st.stop()

# Petit bouton déconnexion en haut à droite
if APP_PASSWORD and st.session_state.get("authenticated"):
    _, col_logout = st.columns([10, 1])
    with col_logout:
        if st.button("🚪 Quitter", use_container_width=True, help="Se déconnecter"):
            st.session_state["authenticated"] = False
            st.rerun()

meta = load_metadata()

# Bandeau état des données
with st.container(border=True):
    st.subheader("📊 État des données")
    c1, c2, c3 = st.columns(3)
    with c1:
        st.metric("Organisations", f"{meta['orgs_rows']:,}".replace(",", " "))
        st.caption(f"Importé : {fmt_date(meta['orgs_uploaded_at'])}")
        st.caption(freshness_status(meta["orgs_uploaded_at"]))
    with c2:
        st.metric("Contacts", f"{meta['contacts_rows']:,}".replace(",", " "))
        st.caption(f"Importé : {fmt_date(meta['contacts_uploaded_at'])}")
        st.caption(freshness_status(meta["contacts_uploaded_at"]))
    with c3:
        success = meta.get("geocoded_success", 0)
        total = meta.get("geocoded_rows", 0)
        rate = f"{success/total*100:.1f}%" if total else "—"
        st.metric("Géocodage BAN", rate)
        st.caption(f"Lancé : {fmt_date(meta['geocoded_at'])}")
        st.caption(f"{success:,} / {total:,} adresses".replace(",", " "))

tab_search, tab_admin = st.tabs(["🔍 Recherche", "⚙️ Admin (import des données)"])

# ─── Onglet Recherche ────────────────────────────────────────────────────────
with tab_search:
    if not os.path.exists(GEOCODED_FILE) or not os.path.exists(CONTACTS_FILE):
        st.warning("⚠️ Aucune donnée chargée. Va dans l'onglet **Admin** pour importer les fichiers Pipedrive.")
    else:
        # Historique (au-dessus du formulaire)
        history = load_history()
        if history:
            with st.expander(f"🕐 Historique des recherches ({len(history)})", expanded=False):
                st.caption("Clique sur une ligne pour rejouer la recherche en un clic.")
                for i, entry in enumerate(history):
                    col_h1, col_h2, col_h3, col_h4 = st.columns([4, 1, 2, 1])
                    with col_h1:
                        st.markdown(f"**{entry['pivot_address']}**")
                        st.caption(entry.get("pivot_label", ""))
                    with col_h2:
                        st.markdown(f"📏 **{entry['radius_m']} m**")
                    with col_h3:
                        st.markdown(
                            f"🏢 {entry['nb_societes']} sociétés  \n"
                            f"✅ {entry['nb_avec_contact']} avec contact"
                        )
                        st.caption(fmt_relative(entry["timestamp"]))
                    with col_h4:
                        if st.button("🔁 Rejouer", key=f"replay_{i}", use_container_width=True):
                            st.session_state["pivot_input"] = entry["pivot_address"]
                            st.session_state["radius_input"] = entry["radius_m"]
                            st.session_state["auto_launch"] = True
                            # Reset le résultat précédent pour forcer une nouvelle recherche
                            st.session_state.pop("last_result", None)
                            st.rerun()
                if st.button("🗑️ Vider l'historique", key="clear_hist"):
                    clear_history()
                    st.rerun()

        # Formulaire de recherche
        # Formulaire de recherche
        col_a, col_b, col_c = st.columns([3, 1, 1])
        with col_a:
            pivot = st.text_input(
                "📍 Adresse pivot",
                value=st.session_state.get("pivot_input", "13 rue d'Uzès, 75002 Paris"),
                placeholder="Ex : 44 avenue des Champs-Élysées, 75008 Paris",
                key="pivot_field",
            )
        with col_b:
            radius = st.number_input(
                "📏 Rayon (m)",
                min_value=100, max_value=5000,
                value=int(st.session_state.get("radius_input", 500)),
                step=100,
                key="radius_field",
            )
        with col_c:
            st.write("")
            st.write("")
            launch_btn = st.button("🚀 Lancer la recherche", type="primary", use_container_width=True)

        # Déclenchement : soit clic explicite, soit replay depuis l'historique
        launch = launch_btn or st.session_state.pop("auto_launch", False)

        # ─── Lancement d'une nouvelle recherche ─────────────────────────────
        if launch and pivot:
            with st.spinner("Recherche en cours…"):
                result, error = run_proximity_search(pivot, radius)
            if error:
                st.error(error)
                st.session_state.pop("last_result", None)
            else:
                # 💾 On stocke le résultat dans session_state pour survivre aux reruns
                st.session_state["last_result"] = result
                st.session_state["last_pivot"] = pivot
                st.session_state["last_radius"] = int(radius)

                # Enregistrer dans l'historique
                add_to_history(
                    pivot_address=pivot,
                    radius_m=int(radius),
                    label=result["pivot_label"],
                    nb_sociétés=len(result["zone"]),
                    nb_avec_contact=len(result["orgs_avec_ids"]),
                )

                # ─── Construction du XLSX (une seule fois, stocké en session) ─
                try:
                    xlsx_bytes = build_export_xlsx(result)
                    filename = f"prospection_{datetime.now().strftime('%Y%m%d_%H%M')}_{int(radius)}m.xlsx"
                    st.session_state["last_xlsx"] = xlsx_bytes
                    st.session_state["last_filename"] = filename
                except Exception as e:
                    import traceback
                    st.error(f"❌ Erreur construction Excel : {type(e).__name__} — {e}")
                    with st.expander("🔍 Détails techniques"):
                        st.code(traceback.format_exc(), language="python")

        # ─── Affichage du dernier résultat (persistant aux reruns) ──────────
        if "last_result" in st.session_state:
            result = st.session_state["last_result"]
            last_radius = st.session_state.get("last_radius", radius)

            st.success(f"📍 Pivot : **{result['pivot_label']}** (score {result['pivot_score']:.2f})")

            zone = result["zone"]
            cz = result["contacts_zone"]
            nb_zone = len(zone)
            nb_avec = len(result["orgs_avec_ids"])
            nb_sans = len(result["orgs_sans_ids"])
            nb_contacts = len(cz)
            nb_tel = int(cz["has_phone"].sum()) if "has_phone" in cz.columns else 0

            # Score global
            st.markdown("### 🎯 Score global")
            m1, m2, m3, m4 = st.columns(4)
            m1.metric("Sociétés zone", nb_zone)
            m2.metric("Avec contact", nb_avec, f"{nb_avec/nb_zone*100:.0f}%" if nb_zone else "—")
            m3.metric("Sans contact", nb_sans, f"{nb_sans/nb_zone*100:.0f}%" if nb_zone else "—")
            m4.metric("Contacts (dont tél.)", nb_contacts,
                      f"{nb_tel} tél." if nb_contacts else "—")

            # ─── Téléchargement (XLSX construit lors du lancement, en cache) ──
            st.markdown("### 📥 Téléchargement")

            if "last_xlsx" in st.session_state and st.session_state["last_xlsx"] is not None:
                st.download_button(
                    "📥 Télécharger le fichier Excel",
                    data=st.session_state["last_xlsx"],
                    file_name=st.session_state.get("last_filename", "prospection.xlsx"),
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    key="dl_main",
                    type="primary",
                )
            else:
                # Fallback : reconstruire à la volée si pas en cache (replay historique)
                try:
                    xlsx_bytes_fb = build_export_xlsx(result)
                    filename_fb = f"prospection_{datetime.now().strftime('%Y%m%d_%H%M')}_{last_radius}m.xlsx"
                    st.download_button(
                        "📥 Télécharger le fichier Excel",
                        data=xlsx_bytes_fb,
                        file_name=filename_fb,
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        key="dl_fallback",
                        type="primary",
                    )
                except Exception as e:
                    import traceback
                    st.error(f"❌ Erreur génération fichier : {type(e).__name__} — {e}")
                    with st.expander("🔍 Détails techniques"):
                        st.code(traceback.format_exc(), language="python")

            # ─── Tableau récap (synthèse de la recherche) ─────────────────
            st.markdown("### 📋 Résumé de la recherche")

            recap_df = pd.DataFrame([
                ["📅 Date de la recherche", datetime.now().strftime("%d/%m/%Y %H:%M")],
                ["📍 Adresse pivot", result["pivot_label"]],
                ["📏 Rayon", f"{last_radius} m"],
                ["🏢 Sociétés dans la zone", str(nb_zone)],
                ["✅ Avec contact", str(nb_avec)],
                ["❌ Sans contact", str(nb_sans)],
                ["👥 Contacts détaillés", str(nb_contacts)],
            ], columns=["Indicateur", "Valeur"])

            st.dataframe(
                recap_df,
                hide_index=True,
                use_container_width=True,
                column_config={
                    "Indicateur": st.column_config.TextColumn(width="medium"),
                    "Valeur": st.column_config.TextColumn(width="large"),
                },
            )

# ─── Onglet Admin ────────────────────────────────────────────────────────────
with tab_admin:
    st.markdown("### 🔐 Zone d'administration")
    pwd = st.text_input("Mot de passe admin", type="password")
    if pwd != ADMIN_PASSWORD:
        if pwd:
            st.error("Mot de passe incorrect.")
        st.info("Saisis le mot de passe pour accéder aux imports.")
    else:
        st.success("✅ Accès admin")

        st.markdown("#### 📊 Import organisations Pipedrive")
        st.caption("Le fichier sera automatiquement géocodé via l'API BAN. Compter ~1-2 min pour 30 000 lignes.")
        orgs_upload = st.file_uploader("Fichier organisations (.xlsx)", type="xlsx", key="orgs_up")
        if orgs_upload and st.button("🚀 Importer et géocoder", key="btn_orgs"):
            try:
                pb = st.progress(0.0, text="Lecture du fichier…")
                file_bytes = orgs_upload.read()
                total, success = process_orgs_file(file_bytes, progress_bar=pb)
                pb.progress(1.0, text="Terminé !")
                now = datetime.now(timezone.utc).isoformat()
                meta["orgs_uploaded_at"] = now
                meta["orgs_rows"] = total
                meta["geocoded_at"] = now
                meta["geocoded_rows"] = total
                meta["geocoded_success"] = success
                save_metadata(meta)
                st.cache_data.clear()
                st.success(f"✅ {total:,} organisations importées, {success:,} géocodées ({success/total*100:.1f}%)".replace(",", " "))
                st.rerun()
            except Exception as e:
                st.error(f"Erreur : {e}")

        st.markdown("---")

        st.markdown("#### 👥 Import contacts Pipedrive")
        st.caption("Doit contenir la colonne 'Organisation - ID' pour le croisement.")

        with st.form("form_contacts", clear_on_submit=False):
            ct_upload = st.file_uploader("Fichier contacts (.xlsx)", type="xlsx", key="ct_up")
            submit_ct = st.form_submit_button("🚀 Importer les contacts", type="primary")

        if submit_ct:
            import traceback
            if not ct_upload:
                st.error("❌ Aucun fichier sélectionné. Glisse le fichier .xlsx puis clique sur Importer.")
            else:
                status_box = st.empty()
                status_box.info(f"📥 Lecture de **{ct_upload.name}** ({ct_upload.size / 1024 / 1024:.1f} Mo)…")
                try:
                    file_bytes = ct_upload.read()
                    status_box.info(f"📦 Fichier lu ({len(file_bytes) / 1024 / 1024:.1f} Mo) — Parsing Excel…")

                    df_test = pd.read_excel(io.BytesIO(file_bytes))
                    status_box.info(f"📊 {len(df_test):,} lignes, {len(df_test.columns)} colonnes — Sauvegarde…".replace(",", " "))

                    if "Organisation - ID" not in df_test.columns:
                        cols_preview = ", ".join(df_test.columns[:10].tolist())
                        st.error(f"❌ Colonne 'Organisation - ID' absente. Colonnes trouvées : {cols_preview}...")
                    else:
                        total = save_contacts_file(file_bytes)
                        meta["contacts_uploaded_at"] = datetime.now(timezone.utc).isoformat()
                        meta["contacts_rows"] = total
                        save_metadata(meta)
                        st.cache_data.clear()
                        status_box.empty()
                        st.success(f"✅ {total:,} contacts importés".replace(",", " "))
                        st.rerun()
                except Exception as e:
                    status_box.empty()
                    st.error(f"❌ Erreur : {type(e).__name__} — {e}")
                    with st.expander("🔍 Détails techniques"):
                        st.code(traceback.format_exc(), language="python")

st.markdown("---")
st.caption("Source géocodage : API officielle Base Adresse Nationale (adresse.data.gouv.fr — IGN/INSEE/DINUM)")
